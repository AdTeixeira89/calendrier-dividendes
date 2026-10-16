"""Lecture des listings : Excel (.xlsx), CSV, PDF texte. Pas d'OCR simulé."""
from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from typing import Any

from .columns import suggest_mapping

MAX_BYTES = 25 * 1024 * 1024
HEADER_SCAN_ROWS = 15


class ReaderError(Exception):
    """Fichier illisible ou format non pris en charge (message affichable)."""


@dataclass
class Table:
    kind: str                      # xlsx | csv | pdf
    headers: list[str]
    rows: list[list[Any]]
    source_rows: list[int]         # numéro de ligne d'origine (1-based) pour chaque ligne
    header_row: int                # index (0-based) dans `raw_rows` de la ligne d'en-tête
    raw_rows: list[list[Any]] = field(default_factory=list)
    sheets: list[str] = field(default_factory=list)
    sheet: str | None = None
    warnings: list[str] = field(default_factory=list)


def detect_kind(filename: str, data: bytes) -> str:
    name = (filename or "").lower()
    if data[:4] == b"PK\x03\x04" or name.endswith((".xlsx", ".xlsm")):
        return "xlsx"
    if data[:5] == b"%PDF-" or name.endswith(".pdf"):
        return "pdf"
    if name.endswith((".csv", ".txt", ".tsv")):
        return "csv"
    if name.endswith(".xls") or data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        raise ReaderError("Format Excel ancien (.xls) non pris en charge : réenregistrer le fichier en .xlsx ou en CSV.")
    raise ReaderError(f"Format de fichier non pris en charge : « {filename} ». Formats acceptés : .xlsx, .csv, .pdf.")


def _cell(v: Any) -> Any:
    if isinstance(v, str):
        return v.strip()
    return v


def _row_nonempty(row: list[Any]) -> bool:
    return any(c is not None and str(c).strip() != "" for c in row)


def _pick_header(raw_rows: list[list[Any]]) -> tuple[int, list[str]]:
    best_i, best_n = 0, -1
    for i, row in enumerate(raw_rows[:HEADER_SCAN_ROWS]):
        if not _row_nonempty(row):
            continue
        n = suggest_mapping([str(c) if c is not None else "" for c in row]).detected
        if n > best_n:
            best_i, best_n = i, n
    row = raw_rows[best_i] if raw_rows else []
    return best_i, _unique_headers(row)


def _unique_headers(row: list[Any]) -> list[str]:
    out: list[str] = []
    seen: dict[str, int] = {}
    for i, c in enumerate(row):
        h = str(c).strip() if c is not None and str(c).strip() else f"Colonne {i + 1}"
        if h in seen:
            seen[h] += 1
            h = f"{h} ({seen[h]})"
        else:
            seen[h] = 1
        out.append(h)
    return out


def _build(kind: str, raw_rows: list[list[Any]], source_numbers: list[int], header_row: int | None,
           warnings: list[str], **extra: Any) -> Table:
    if not any(_row_nonempty(r) for r in raw_rows):
        raise ReaderError("Le fichier ne contient aucune donnée.")
    if header_row is None:
        header_row, headers = _pick_header(raw_rows)
        if suggest_mapping(headers).detected < 2:
            warnings.append("Ligne d'en-tête non reconnue automatiquement : vérifier la ligne d'en-tête et la correspondance des colonnes.")
    else:
        if not (0 <= header_row < len(raw_rows)):
            raise ReaderError(f"Ligne d'en-tête invalide : {header_row}.")
        headers = _unique_headers(raw_rows[header_row])
    width = len(headers)
    rows, nums = [], []
    for r, n in zip(raw_rows[header_row + 1:], source_numbers[header_row + 1:]):
        r = list(r) + [None] * (width - len(r))
        rows.append(r[:width] if len(r) > width and not _row_nonempty(r[width:]) else r)
        nums.append(n)
    # étendre les en-têtes si des données dépassent
    maxw = max([len(r) for r in rows] + [width])
    if maxw > width:
        headers += [f"Colonne {i + 1}" for i in range(width, maxw)]
        rows = [r + [None] * (maxw - len(r)) for r in rows]
    return Table(kind=kind, headers=headers, rows=rows, source_rows=nums, header_row=header_row,
                 raw_rows=raw_rows, warnings=warnings, **extra)


def read_xlsx(data: bytes, sheet: str | None = None, header_row: int | None = None) -> Table:
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ReaderError(f"Fichier Excel illisible ou corrompu : {exc}") from exc
    names = wb.sheetnames
    warnings: list[str] = []
    if sheet is None:
        sheet = next((n for n in names if _sheet_has_data(wb[n])), names[0] if names else None)
        if len(names) > 1:
            warnings.append(f"{len(names)} feuilles dans le classeur ; feuille « {sheet} » utilisée (modifiable).")
    if sheet not in names:
        raise ReaderError(f"Feuille introuvable : « {sheet} ».")
    raw, nums = [], []
    for i, row in enumerate(wb[sheet].iter_rows(values_only=True), start=1):
        cells = [_cell(c) for c in row]
        while cells and (cells[-1] is None or cells[-1] == ""):
            cells.pop()
        raw.append(cells)
        nums.append(i)
    wb.close()
    return _build("xlsx", raw, nums, header_row, warnings, sheets=names, sheet=sheet)


def _sheet_has_data(ws) -> bool:
    for row in ws.iter_rows(values_only=True):
        if _row_nonempty(list(row)):
            return True
    return False


def _decode(data: bytes) -> tuple[str, str | None]:
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(enc), (None if enc == "utf-8-sig" else "Encodage Windows-1252 supposé (fichier non UTF-8).")
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1"), "Encodage Latin-1 supposé."


def read_csv(data: bytes, header_row: int | None = None) -> Table:
    text, enc_warn = _decode(data)
    warnings = [enc_warn] if enc_warn else []
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
    except csv.Error:
        first = sample.splitlines()[0] if sample.splitlines() else ""
        delim = max(";,\t|", key=first.count)
        dialect = csv.excel
        dialect.delimiter = delim  # type: ignore[misc]
        warnings.append(f"Séparateur CSV deviné : « {delim} ».")
    reader = csv.reader(io.StringIO(text), dialect)
    raw, nums = [], []
    try:
        for row in reader:
            raw.append([_cell(c) for c in row])
            nums.append(reader.line_num)
    except csv.Error as exc:
        raise ReaderError(f"CSV invalide ligne {reader.line_num} : {exc}") from exc
    return _build("csv", raw, nums, header_row, warnings)


def read_pdf(data: bytes, header_row: int | None = None) -> Table:
    import pdfplumber
    try:
        pdf = pdfplumber.open(io.BytesIO(data))
    except Exception as exc:
        raise ReaderError(f"PDF illisible ou protégé : {exc}") from exc
    warnings: list[str] = []
    table_rows: list[list[Any]] = []
    text_lines: list[str] = []
    with pdf:
        for page in pdf.pages:
            for tbl in page.extract_tables() or []:
                for r in tbl:
                    cells = [(c or "").replace("\n", " ").strip() for c in r]
                    if _row_nonempty(cells):
                        table_rows.append(cells)
            text_lines += (page.extract_text() or "").splitlines()
    if len("".join(text_lines).strip()) < 20 and not table_rows:
        raise ReaderError(
            "PDF sans texte exploitable (probablement un scan). La lecture par OCR n'est pas disponible dans cette version : "
            "fournir le listing en Excel/CSV ou un PDF texte."
        )
    if table_rows:
        header_candidate = table_rows[0]
        # supprime les en-têtes répétés à chaque page
        raw = [table_rows[0]] + [r for r in table_rows[1:] if r != header_candidate]
        warnings.append("Tableau extrait d'un PDF : vérifier les colonnes et les valeurs avant import.")
    else:
        import re
        raw = [re.split(r"\s{2,}|\t|;", ln.strip()) for ln in text_lines if ln.strip()]
        warnings.append("PDF sans tableau détecté : lignes découpées d'après les espaces. Résultat à vérifier attentivement.")
    return _build("pdf", raw, list(range(1, len(raw) + 1)), header_row, warnings)


def read_listing(filename: str, data: bytes, sheet: str | None = None, header_row: int | None = None) -> Table:
    if not data:
        raise ReaderError("Fichier vide.")
    if len(data) > MAX_BYTES:
        raise ReaderError(f"Fichier trop volumineux (> {MAX_BYTES // (1024 * 1024)} Mo).")
    kind = detect_kind(filename, data)
    if kind == "xlsx":
        return read_xlsx(data, sheet, header_row)
    if kind == "pdf":
        return read_pdf(data, header_row)
    return read_csv(data, header_row)
